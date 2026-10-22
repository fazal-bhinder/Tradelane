"""HTTP API. Every evaluation is computed on read from the rules, so editing a
rate or a product updates every screen at once; a review stores the snapshot
it was decided on, and every change lands in the activity log with its author."""

from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
import os
import time
from dataclasses import asdict
from pathlib import Path

from fastapi import Depends, FastAPI, HTTPException, Request, Response, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import delete, select, text
from sqlalchemy.orm import Session

from . import ai, analyst
from .catalogue import latest_reviews, lanes_for, stats, to_engine, world
from .db import Ask, Cache, Event, ProductRow, Review, Setting, User, init_db, now, session
from .engine import World, evaluate, sensitivity
from .rules import ASSUMPTION_FIELDS, CATEGORIES, CHANNELS, CHECKED, DUTIES, ORIGINS, ROUTES, US_MFN_LINES
from .sources import us_mfn

app = FastAPI(title="Tradelane", version="1.1")
SAMPLE = Path(__file__).resolve().parent.parent / "data" / "sample_catalogue.csv"
SECRET = os.environ.get("SESSION_SECRET", "tradelane-demo-session").encode()
COOKIE = "tl_session"
_ready = False


def db():
    global _ready
    if not _ready:
        init_db()
        _ready = True
    s = session()
    try:
        yield s
    finally:
        s.close()


# ---- demo sign-in: one click per person, a signed cookie, no passwords ----

def _sign(user_id: str) -> str:
    return f"{user_id}.{hmac.new(SECRET, user_id.encode(), hashlib.sha256).hexdigest()[:32]}"


def current_user(request: Request, s: Session = Depends(db)) -> User:
    raw = request.cookies.get(COOKIE, "")
    uid = raw.split(".", 1)[0]
    if not raw or not hmac.compare_digest(raw, _sign(uid)) or not (u := s.get(User, uid)):
        raise HTTPException(401, "Sign in to continue")
    return u


def user_json(u: User) -> dict:
    return {"id": u.id, "name": u.name, "role": u.role, "title": u.title, "color": u.color}


def log(s: Session, u: User | None, kind: str, message: str, product: ProductRow | None = None,
        channel: str | None = None, **data) -> None:
    s.add(Event(user_id=u.id if u else None, kind=kind, message=message, product_id=product.id if product else None,
                channel=channel, data=data))


@app.get("/api/users")
def users(s: Session = Depends(db)):
    return [user_json(u) for u in s.scalars(select(User).order_by(User.id))]


class LoginIn(BaseModel):
    user_id: str


@app.post("/api/login")
def login(body: LoginIn, response: Response, s: Session = Depends(db)):
    u = s.get(User, body.user_id) or _404("No such demo user")
    response.set_cookie(COOKIE, _sign(u.id), httponly=True, samesite="lax", secure=bool(os.environ.get("VERCEL")),
                        max_age=60 * 60 * 24 * 14)
    log(s, u, "login", f"{u.name.split()[0]} signed in")
    s.commit()
    return user_json(u)


@app.post("/api/logout")
def logout(response: Response):
    response.delete_cookie(COOKIE)
    return {"ok": True}


@app.get("/api/me")
def me(u: User = Depends(current_user)):
    return user_json(u)


@app.get("/api/health")
def health(s: Session = Depends(db)):
    t = time.perf_counter()
    s.execute(text("select 1"))
    ms = round((time.perf_counter() - t) * 1000)
    return {"ok": True, "ai": ai.enabled(), "db_ms": ms}


# ---- products ----

def result_json(res) -> dict:
    d = asdict(res)
    d["lines"] = [{**asdict(l)} for l in res.lines]
    return d


def product_json(r: ProductRow, w: World, reviews: dict[tuple[int, str], Review], full: bool = False) -> dict:
    p = to_engine(r)
    lanes = []
    for cid in lanes_for(r):
        res = evaluate(p, cid, w)
        rv = reviews.get((r.id, cid))
        lane = result_json(res) if full else {
            k: getattr(res, k) for k in ("channel", "lane", "currency", "price", "profit", "margin", "break_even",
                                         "target_price", "verdict", "reasons", "missing", "net_revenue")}
        if not full:
            lane["costs"] = {l.key: round(l.amount, 4) for l in res.lines}
        lane["review"] = {"decision": rv.decision, "note": rv.note, "at": rv.created_at.isoformat(), "by": rv.user_id} if rv else None
        if full:
            lane["sensitivity"] = [asdict(x) for x in sensitivity(p, cid, w)]
        lanes.append(lane)
    return {
        "id": r.id, "sku": r.sku, "name": r.name, "description": r.description, "supplier": r.supplier,
        "unit_cost": r.unit_cost, "cost_currency": r.cost_currency, "weight_kg": r.weight_kg,
        "dims_cm": [r.length_cm, r.width_cm, r.height_cm], "origin": r.origin, "category": r.category,
        "category_kind": r.category_kind, "hs_code": r.hs_code, "us_mfn": r.us_mfn, "us_mfn_line": r.us_mfn_line,
        "prices": r.prices or {}, "ai": r.ai or {}, "lanes": lanes,
    }


@app.get("/api/rules")
def rules(s: Session = Depends(db), _: User = Depends(current_user)):
    w = world(s)
    overrides = {r.key: r.value for r in s.scalars(select(Setting))}
    return {
        "checked": CHECKED,
        "categories": CATEGORIES,
        "origins": ORIGINS,
        "fx": {"rates": w.fx, "source": asdict(w.fx_source)},
        "assumptions": [{"key": k, "label": lbl, "unit": unit, "value": getattr(w.assumptions, k), "default": d,
                         "edited": k in overrides} for k, (lbl, unit, d) in ASSUMPTION_FIELDS.items()],
        "routes": [{"origin": o, "dest": d, "currency": r.currency, "base": r.base, "per_kg": r.per_kg,
                    "transit_days": r.transit_days, "source": asdict(r.source)} for (o, d), r in w.routes.items()],
        "channels": [{"id": c.id, "name": c.name, "country": c.country, "currency": c.currency,
                      "vat_rate": c.vat_rate, "vat_source": asdict(c.vat_source),
                      "fees": {cat: {"tiers": [asdict(t) for t in f.tiers], "whole_price": f.whole_price,
                                     "min_fee": f.min_fee, "source": asdict(f.source)} for cat, f in c.fees.items()}}
                     for c in w.channels.values()],
        "duties": {cc: {"category_rates": d.category_rates, "extra_by_origin": d.extra_by_origin,
                        "relief_threshold": d.relief_threshold, "basis": "CIF" if d.includes_shipping else "FOB",
                        "source": asdict(d.source), "extra_source": asdict(d.extra_source)}
                   for cc, d in DUTIES.items()},
        "us_mfn_lines": {k: {"line": a, "what": b, "kind": c} for k, (a, b, c) in US_MFN_LINES.items()},
    }


class SettingIn(BaseModel):
    key: str
    value: float = Field(ge=0)


def _setting_label(key: str) -> str:
    if key in ASSUMPTION_FIELDS:
        return ASSUMPTION_FIELDS[key][0].lower()
    _, lane, field = key.split(":")
    o, d = lane.split("-")
    return f"{ORIGINS[o]} to {'UK' if d == 'GB' else 'US'} shipping {'per kg' if field == 'per_kg' else 'per parcel'}"


@app.put("/api/rules")
def update_rule(body: SettingIn, s: Session = Depends(db), u: User = Depends(current_user)):
    valid = set(ASSUMPTION_FIELDS) | {f"route:{o}-{d}:{f}" for (o, d) in ROUTES for f in ("base", "per_kg")}
    if body.key not in valid:
        raise HTTPException(422, f"Unknown setting {body.key}")
    if body.key in ASSUMPTION_FIELDS and body.value >= 1:
        raise HTTPException(422, "Percentages are fractions: 0.2 means 20%")
    s.merge(Setting(key=body.key, value=body.value))
    shown = f"{body.value:.1%}" if body.key in ASSUMPTION_FIELDS else f"{body.value:g}"
    log(s, u, "rule", f"set {_setting_label(body.key)} to {shown}")
    s.commit()
    return {"ok": True}


@app.delete("/api/rules/{key}")
def reset_rule(key: str, s: Session = Depends(db), u: User = Depends(current_user)):
    if s.execute(delete(Setting).where(Setting.key == key)).rowcount:
        log(s, u, "rule", f"reset {_setting_label(key)} to the default")
    s.commit()
    return {"ok": True}


@app.get("/api/products")
def list_products(s: Session = Depends(db), _: User = Depends(current_user)):
    w, reviews = world(s), latest_reviews(s)
    rows = s.scalars(select(ProductRow).order_by(ProductRow.id)).all()
    return {"products": [product_json(r, w, reviews) for r in rows],
            "fx": {"rates": w.fx, "source": asdict(w.fx_source)},
            "target_margin": w.assumptions.target_margin, "ai": ai.enabled()}


@app.get("/api/products/{pid}")
def get_product(pid: int, s: Session = Depends(db), _: User | None = Depends(current_user)):
    r = s.get(ProductRow, pid) or _404()
    history = s.scalars(select(Review).where(Review.product_id == pid).order_by(Review.created_at.desc())).all()
    out = product_json(r, world(s), latest_reviews(s), full=True)
    out["history"] = [{"channel": h.channel, "decision": h.decision, "note": h.note, "at": h.created_at.isoformat(), "by": h.user_id}
                      for h in history]
    return out


class ProductPatch(BaseModel):
    name: str | None = None
    unit_cost: float | None = Field(None, gt=0)
    weight_kg: float | None = Field(None, gt=0)
    category: str | None = None
    hs_code: str | None = None
    prices: dict[str, float | None] | None = None


@app.patch("/api/products/{pid}")
def patch_product(pid: int, body: ProductPatch, s: Session = Depends(db), u: User = Depends(current_user)):
    r = s.get(ProductRow, pid) or _404()
    data = body.model_dump(exclude_unset=True)
    if "category" in data:
        if data["category"] not in CATEGORIES:
            raise HTTPException(422, f"Category must be one of {list(CATEGORIES)}")
        was = r.category_kind
        r.category, r.category_kind = data.pop("category"), "fact"  # a person confirmed or chose it
        log(s, u, "category", f"{'confirmed' if was in ('ai', 'assumption') else 'set'} the category as {CATEGORIES[r.category].lower()}", r)
    if "hs_code" in data:
        r.hs_code = data.pop("hs_code") or None
        _apply_hs(s, r)
        log(s, u, "hs", f"set HS code {r.hs_code or 'none'}" + (f", US duty {r.us_mfn:.1%}" if r.us_mfn is not None else ""), r)
    if "prices" in data:
        changed = {k: v for k, v in data.pop("prices").items() if k in CHANNELS}
        r.prices = {**(r.prices or {}), **changed}
        for k, v in changed.items():
            cur = CHANNELS[k].currency
            log(s, u, "price", f"priced it at {cur} {v:,.2f}" if v is not None else "cleared the price", r, k, price=v)
    for k, v in data.items():
        setattr(r, k, v)
        log(s, u, "edit", f"changed {k.replace('_', ' ')} to {v}", r)
    s.commit()
    return get_product(pid, s)


def _apply_hs(s: Session, r: ProductRow) -> None:
    hit = us_mfn(s, r.hs_code) if r.hs_code else None
    r.us_mfn, r.us_mfn_line = (hit if hit else (None, None))


@app.post("/api/products/{pid}/classify")
def classify(pid: int, s: Session = Depends(db), u: User = Depends(current_user)):
    r = s.get(ProductRow, pid) or _404()
    t = time.perf_counter()
    try:
        sug, by = ai.classify(r.name, r.description, r.category)
    except Exception as e:  # the model or gateway failed: say so, change nothing
        raise HTTPException(502, f"Classification failed: {e}") from e
    ms = round((time.perf_counter() - t) * 1000)
    r.ai = {**sug.model_dump(), "by": by, "ms": ms}
    if r.category_kind != "fact":
        r.category, r.category_kind = sug.category, "ai" if by != "keyword match" else "assumption"
    if not r.hs_code:
        r.hs_code = sug.hs_code
        _apply_hs(s, r)
    who = "matched it by keyword" if by == "keyword match" else "had the AI classify it"
    log(s, u, "classify", f"{who} as {CATEGORIES[sug.category].lower()}, HS {sug.hs_code} ({sug.confidence:.0%} sure)", r, by=by, ms=ms)
    s.commit()
    return get_product(pid, s)


@app.delete("/api/products/{pid}")
def delete_product(pid: int, s: Session = Depends(db), u: User = Depends(current_user)):
    r = s.get(ProductRow, pid) or _404()
    log(s, u, "delete", f"deleted {r.name}")
    s.execute(delete(Review).where(Review.product_id == pid))
    s.execute(delete(ProductRow).where(ProductRow.id == pid))
    s.commit()
    return {"ok": True}


COLUMNS = ["sku", "name", "description", "supplier", "unit_cost", "cost_currency", "weight_kg", "length_cm",
           "width_cm", "height_cm", "origin", "category", "hs_code", "price_amazon_uk", "price_amazon_us",
           "price_walmart_us"]


def _num(v) -> float | None:
    if v in (None, ""):
        return None
    return float(str(v).replace(",", "").strip())


def parse_rows(rows: list[dict]) -> tuple[list[ProductRow], list[str]]:
    out, errors = [], []
    for i, raw in enumerate(rows, start=2):  # row 1 is the header
        row = {(k or "").strip().lower(): (v.strip() if isinstance(v, str) else v) for k, v in raw.items()}
        if not any(row.values()):
            continue
        try:
            if not row.get("name"):
                raise ValueError("name is empty")
            origin = (row.get("origin") or "IN").upper()
            if origin not in ORIGINS:
                raise ValueError(f"origin must be one of {', '.join(ORIGINS)}")
            cat = (row.get("category") or "").lower() or None
            if cat and cat not in CATEGORIES:
                raise ValueError(f"category '{cat}' is not one of {', '.join(CATEGORIES)}")
            out.append(ProductRow(
                sku=row.get("sku") or f"ROW-{i}", name=row["name"], description=row.get("description") or "",
                supplier=row.get("supplier") or "", unit_cost=_num(row.get("unit_cost")),
                cost_currency=(row.get("cost_currency") or ("INR" if origin == "IN" else "GBP")).upper(),
                weight_kg=_num(row.get("weight_kg")), length_cm=_num(row.get("length_cm")),
                width_cm=_num(row.get("width_cm")), height_cm=_num(row.get("height_cm")), origin=origin,
                category=cat, category_kind="fact" if cat else "unknown", hs_code=row.get("hs_code") or None,
                prices={c: _num(row.get(f"price_{c}")) for c in CHANNELS if _num(row.get(f"price_{c}")) is not None},
            ))
        except (ValueError, TypeError) as e:
            errors.append(f"Row {i}: {e}")
    return out, errors


def _read_upload(name: str, data: bytes) -> list[dict]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "") for h in next(it)]
        return [dict(zip(header, r)) for r in it]
    return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))


def _save(s: Session, rows: list[ProductRow], replace_all: bool) -> None:
    if replace_all:
        s.execute(delete(Review))
        s.execute(delete(ProductRow))
    s.add_all(rows)
    s.flush()
    for r in rows:
        if r.hs_code:
            _apply_hs(s, r)


@app.post("/api/products/import")
async def import_products(file: UploadFile, replace_all: bool = False, s: Session = Depends(db), u: User = Depends(current_user)):
    if not file.filename or not file.filename.lower().endswith((".csv", ".xlsx", ".xlsm")):
        raise HTTPException(422, "Upload a .csv or .xlsx file")
    data = await file.read()
    if len(data) > 4_000_000:
        raise HTTPException(413, "Keep the file under 4 MB")
    try:
        rows, errors = parse_rows(_read_upload(file.filename, data))
    except (UnicodeDecodeError, StopIteration, csv.Error) as e:
        raise HTTPException(422, f"Could not read the file: {e}") from e
    if not rows:
        raise HTTPException(422, "No product rows found. " + " ".join(errors[:5]))
    _save(s, rows, replace_all)
    log(s, u, "import", f"imported {len(rows)} products from {file.filename}" + (f", skipped {len(errors)} rows" if errors else ""))
    s.commit()
    return {"imported": len(rows), "errors": errors}


@app.post("/api/products/sample")
def load_sample(s: Session = Depends(db), u: User = Depends(current_user)):
    rows, errors = parse_rows(list(csv.DictReader(SAMPLE.open(encoding="utf-8"))))
    _save(s, rows, replace_all=True)
    log(s, u, "import", f"loaded the {len(rows)}-product sample catalogue")
    s.commit()
    return {"imported": len(rows), "errors": errors}


@app.get("/api/template.csv")
def template():
    return Response(",".join(COLUMNS) + "\n", media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=tradelane-template.csv"})


# ---- review and export ----

class ReviewIn(BaseModel):
    product_id: int
    channel: str
    decision: str = Field(pattern="^(approved|rejected)$")
    note: str = Field("", max_length=500)


@app.post("/api/reviews")
def add_review(body: ReviewIn, s: Session = Depends(db), u: User = Depends(current_user)):
    if u.role != "approver":
        raise HTTPException(403, "Only an approver can decide a lane. Sign in as Rohan to approve.")
    r = s.get(ProductRow, body.product_id) or _404()
    if body.channel not in lanes_for(r):
        raise HTTPException(422, "That lane does not exist for this product")
    res = evaluate(to_engine(r), body.channel, world(s))
    if body.decision == "approved" and res.verdict == "needs-data":
        raise HTTPException(422, "Fill the missing data before approving: " + ", ".join(res.missing))
    s.add(Review(product_id=r.id, channel=body.channel, decision=body.decision, note=body.note,
                 snapshot=result_json(res), user_id=u.id))
    margin = f" at {res.margin:.1%} margin" if res.margin is not None else ""
    log(s, u, "review", f"{body.decision} {CHANNELS[body.channel].name}{margin}" + (f": “{body.note}”" if body.note else ""),
        r, body.channel, decision=body.decision)
    s.commit()
    return {"ok": True}


@app.get("/api/export.csv")
def export(channel: str, s: Session = Depends(db), _: User = Depends(current_user)):
    if channel not in CHANNELS:
        raise HTTPException(404, "Unknown channel")
    reviews = latest_reviews(s)
    buf = io.StringIO()
    out = csv.writer(buf)
    out.writerow(["sku", "title", "bullet_1", "bullet_2", "bullet_3", "category", "hs_code", "price",
                  "currency", "net_profit", "margin", "approved_at", "approved_by", "note"])
    for r in s.scalars(select(ProductRow).order_by(ProductRow.id)):
        rv = reviews.get((r.id, channel))
        if not rv or rv.decision != "approved":
            continue
        snap = rv.snapshot
        bullets = (r.ai or {}).get("bullets", []) + ["", "", ""]
        out.writerow([r.sku, (r.ai or {}).get("title") or r.name, *bullets[:3], r.category, r.hs_code or "",
                      snap["price"], snap["currency"], snap["profit"],
                      f"{snap['margin']:.4f}" if snap["margin"] is not None else "", rv.created_at.date().isoformat(),
                      rv.user_id or "", rv.note])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=tradelane-{channel}-approved.csv"})


# ---- activity, analyst, brief ----

@app.get("/api/activity")
def activity(limit: int = 30, s: Session = Depends(db), _: User = Depends(current_user)):
    names = {p.id: p.name for p in s.scalars(select(ProductRow))}
    rows = s.scalars(select(Event).where(Event.kind != "login").order_by(Event.created_at.desc()).limit(min(limit, 100))).all()
    return [{"id": e.id, "user": e.user_id, "kind": e.kind, "message": e.message, "product_id": e.product_id,
             "product": names.get(e.product_id), "channel": e.channel, "at": e.created_at.isoformat()} for e in rows]


class AskIn(BaseModel):
    question: str = Field(min_length=3, max_length=500)
    sku: str | None = None


@app.post("/api/ask")
def ask(body: AskIn, s: Session = Depends(db), u: User = Depends(current_user)):
    log(s, u, "ask", f"asked the analyst: “{body.question}”")
    s.commit()

    uid = u.id

    def stream():
        # its own session: the request's may be closed before a long answer finishes streaming
        with session() as s2:
            for event in analyst.ask(s2, uid, body.question, body.sku):
                yield f"data: {json.dumps(event, default=str)}\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/asks")
def asks(s: Session = Depends(db), u: User = Depends(current_user)):
    rows = s.scalars(select(Ask).order_by(Ask.created_at.desc()).limit(12)).all()
    return [{"id": a.id, "user": a.user_id, "question": a.question, "answer": a.answer, "steps": a.steps, "ms": a.ms,
             "at": a.created_at.isoformat()} for a in rows]


@app.get("/api/brief")
def get_brief(s: Session = Depends(db), _: User = Depends(current_user)):
    row = s.get(Cache, "brief")
    return row.value | {"at": row.fetched_at.isoformat()} if row else None


@app.post("/api/brief")
def make_brief(s: Session = Depends(db), u: User = Depends(current_user)):
    t = time.perf_counter()
    try:
        text_, by = analyst.brief(s)
    except Exception as e:
        raise HTTPException(502, f"The brief failed: {e}") from e
    value = {"text": text_, "by": by, "ms": round((time.perf_counter() - t) * 1000), "for": u.id}
    s.merge(Cache(key="brief", value=value, fetched_at=now()))
    log(s, u, "brief", "had the AI write the brief" if by != "template" else "wrote the brief from a template")
    s.commit()
    return value | {"at": now().isoformat()}


@app.get("/api/stats")
def get_stats(s: Session = Depends(db), _: User = Depends(current_user)):
    return stats(s)


def _404(msg: str = "Product not found"):
    raise HTTPException(404, msg)
